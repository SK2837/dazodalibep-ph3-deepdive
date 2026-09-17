"""Shared helpers for building the dazodalibep Phase 3 deep-dive deck."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

# ---------------------------------------------------------------- palette --
NAVY = RGBColor(0x1B, 0x2A, 0x4A)
TEAL = RGBColor(0x2E, 0x7D, 0x7B)
GRAY = RGBColor(0x8C, 0x8C, 0x8C)
LIGHT = RGBColor(0xF2, 0xF2, 0xF0)
RED = RGBColor(0xB4, 0x45, 0x3E)
GOLD = RGBColor(0xB0, 0x8A, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x22, 0x22, 0x22)
RED_TINT = RGBColor(0xF7, 0xE8, 0xE7)
GOLD_TINT = RGBColor(0xF6, 0xEE, 0xD6)
TEAL_TINT = RGBColor(0xE6, 0xF1, 0xF0)
GRAY_TINT = RGBColor(0xF4, 0xF4, 0xF3)

FONT = "Calibri"

PAGE_W = 13.333
PAGE_H = 7.5
L = 0.55
R = 0.55
CW = PAGE_W - L - R
TITLE_TOP = 0.38
BODY_TOP = 1.55
FOOTER_TOP = 6.98

# Standard two-column split used by several slides (left content + right panel)
LEFT_W = 6.25
COL_GAP = 0.35
RIGHT_X = L + LEFT_W + COL_GAP
RIGHT_W = PAGE_W - R - RIGHT_X

# ------------------------------------------------------------- utilities --

def new_deck():
    prs = Presentation()
    prs.slide_width = Emu(int(PAGE_W * 914400))
    prs.slide_height = Emu(int(PAGE_H * 914400))
    return prs


def blank_slide(prs):
    layout = prs.slide_layouts[6]
    return prs.slides.add_slide(layout)


def set_bg(slide, color):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color


def set_alpha(shape, pct):
    """pct = opacity percent (0-100). Applies to solid fill."""
    sp = shape.fill.fore_color._xFill
    srgb = sp.find(qn('a:srgbClr'))
    if srgb is None:
        return
    alpha = srgb.makeelement(qn('a:alpha'), {'val': str(int(pct * 1000))})
    srgb.append(alpha)


def add_text(slide, text, left, top, width, height, size=18, color=DARK,
             bold=False, italic=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             font=FONT, line_spacing=1.0, wrap=True, shrink=False):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = font
        r.font.color.rgb = color
    return box


def add_title(slide, text, color=NAVY, size=None):
    if size is None:
        n = len(text)
        size = 30 if n <= 78 else (28 if n <= 105 else 25)
    return add_text(slide, text, L, TITLE_TOP, CW, 1.12, size=size, color=color,
                     bold=True, anchor=MSO_ANCHOR.TOP, line_spacing=1.02)


def add_footer(slide, source_text, page_no):
    add_text(slide, "Source: " + source_text, L, FOOTER_TOP, CW - 1.0, 0.4,
              size=10, color=GRAY, italic=True, anchor=MSO_ANCHOR.TOP)
    add_text(slide, page_no, PAGE_W - R - 0.6, FOOTER_TOP, 0.6, 0.4,
              size=10, color=GRAY, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.TOP)


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def add_bullets(slide, items, left, top, width, height, size=18, color=DARK,
                 gap_after=6, bullet_char="•", lead_bold_split=None,
                 line_spacing=1.08):
    """items: list of strings (or (lead, rest) tuples if lead_bold_split)."""
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.line_spacing = line_spacing
        p.space_after = Pt(gap_after)
        if isinstance(item, tuple):
            lead, rest = item
            r1 = p.add_run()
            r1.text = f"{bullet_char}  {lead}"
            r1.font.bold = True
            r1.font.size = Pt(size)
            r1.font.name = FONT
            r1.font.color.rgb = color
            r2 = p.add_run()
            r2.text = rest
            r2.font.bold = False
            r2.font.size = Pt(size)
            r2.font.name = FONT
            r2.font.color.rgb = color
        else:
            r = p.add_run()
            r.text = f"{bullet_char}  {item}" if bullet_char else item
            r.font.size = Pt(size)
            r.font.name = FONT
            r.font.color.rgb = color
    return box


def add_numbered(slide, items, left, top, width, height, size=16, color=DARK,
                  gap_after=8, line_spacing=1.08):
    """items: list of (lead_bold, rest) tuples, auto-numbered 1.,2.,3."""
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    first = True
    for i, (lead, rest) in enumerate(items, start=1):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.line_spacing = line_spacing
        p.space_after = Pt(gap_after)
        r0 = p.add_run()
        r0.text = f"{i}.  "
        r0.font.bold = True
        r0.font.size = Pt(size)
        r0.font.name = FONT
        r0.font.color.rgb = NAVY
        r1 = p.add_run()
        r1.text = lead
        r1.font.bold = True
        r1.font.size = Pt(size)
        r1.font.name = FONT
        r1.font.color.rgb = color
        r2 = p.add_run()
        r2.text = rest
        r2.font.bold = False
        r2.font.size = Pt(size)
        r2.font.name = FONT
        r2.font.color.rgb = color
    return box


def callout_box(slide, left, top, width, height, label, body, fill=RED_TINT,
                 border=RED, label_color=RED, body_color=DARK, label_size=12,
                 body_size=14):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top),
                                  Inches(width), Inches(height))
    shp.adjustments[0] = 0.06
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = border
    shp.line.width = Pt(1)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.18)
    tf.margin_top = Inches(0.1)
    tf.margin_bottom = Inches(0.1)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p0 = tf.paragraphs[0]
    r0 = p0.add_run()
    r0.text = label
    r0.font.bold = True
    r0.font.size = Pt(label_size)
    r0.font.name = FONT
    r0.font.color.rgb = label_color
    p1 = tf.add_paragraph()
    p1.space_before = Pt(3)
    r1 = p1.add_run()
    r1.text = body
    r1.font.size = Pt(body_size)
    r1.font.name = FONT
    r1.font.color.rgb = body_color
    return shp


def card_box(slide, left, top, width, height, fill=LIGHT, border=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top),
                                  Inches(width), Inches(height))
    shp.adjustments[0] = 0.045
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if border:
        shp.line.color.rgb = border
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_picture_fit(slide, path, left, top, max_w, max_h, align="center"):
    from PIL import Image
    im = Image.open(path)
    iw, ih = im.size
    ar = iw / ih
    w = max_w
    h = w / ar
    if h > max_h:
        h = max_h
        w = h * ar
    if align == "center":
        left = left + (max_w - w) / 2
    elif align == "right":
        left = left + (max_w - w)
    return slide.shapes.add_picture(path, Inches(left), Inches(top), Inches(w), Inches(h))


def _shade_cell(cell, color):
    cell.fill.solid()
    cell.fill.fore_color.rgb = color


def add_table(slide, data, left, top, width, height, col_widths=None,
              row_heights=None, header=True, header_fill=NAVY, header_color=WHITE,
              body_color=DARK, font_size=12, header_size=None, align_cols=None,
              bold_col0=False, cell_shades=None, valign=MSO_ANCHOR.MIDDLE,
              banding=None):
    """data: list of rows, each a list of cell strings.
    align_cols: dict col_idx -> PP_ALIGN
    cell_shades: dict (r,c) -> RGBColor override
    banding: RGBColor for alternating body rows (applied where no override)
    """
    rows = len(data)
    cols = len(data[0])
    header_size = header_size or font_size
    shape = slide.shapes.add_table(rows, cols, Inches(left), Inches(top),
                                    Inches(width), Inches(height))
    table = shape.table
    # kill python-pptx default banding theme so our fills show cleanly
    tbl = table._tbl
    tblPr = tbl.find(qn('a:tblPr'))
    if tblPr is not None:
        tblPr.set('firstRow', '0')
        tblPr.set('bandRow', '0')
    if col_widths:
        total = sum(col_widths)
        for i, w in enumerate(col_widths):
            table.columns[i].width = Inches(width * w / total)
    if row_heights:
        for i, h in enumerate(row_heights):
            table.rows[i].height = Inches(h)
    else:
        for i in range(rows):
            table.rows[i].height = Inches(height / rows)
    align_cols = align_cols or {}
    cell_shades = cell_shades or {}
    for r in range(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = valign
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = align_cols.get(c, PP_ALIGN.LEFT)
            text = str(data[r][c])
            is_header = header and r == 0
            first_line = True
            for line in text.split("\n"):
                pp = p if first_line else tf.add_paragraph()
                pp.alignment = align_cols.get(c, PP_ALIGN.LEFT)
                first_line = False
                run = pp.add_run()
                run.text = line
                run.font.name = FONT
                run.font.size = Pt(header_size if is_header else font_size)
                run.font.bold = is_header or (bold_col0 and c == 0)
                run.font.color.rgb = header_color if is_header else body_color
            if is_header:
                _shade_cell(cell, header_fill)
            elif (r, c) in cell_shades:
                _shade_cell(cell, cell_shades[(r, c)])
            elif banding and r % 2 == 0:
                _shade_cell(cell, banding)
            else:
                _shade_cell(cell, WHITE)
    return shape
